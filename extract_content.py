#!/usr/bin/env python3
"""Extract full content of the Word document to understand its structure."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

doc = Document("/home/melquisedec/Escritorio/Projects/Progra III/Proyecto Final/Project Document.docx")

print("=" * 80)
print("DOCUMENT STRUCTURE ANALYSIS")
print("=" * 80)

# Print all paragraphs with their styles and formatting details
for i, para in enumerate(doc.paragraphs):
    style = para.style.name if para.style else "None"
    text = para.text.strip()
    if text:
        # Get alignment
        align = str(para.alignment) if para.alignment else "None"
        # Get font info from runs
        fonts = []
        sizes = []
        bolds = []
        for run in para.runs:
            if run.font.name:
                fonts.append(run.font.name)
            if run.font.size:
                sizes.append(str(run.font.size))
            bolds.append(str(run.font.bold))
        
        print(f"\n--- Paragraph {i} ---")
        print(f"  Style: {style}")
        print(f"  Alignment: {align}")
        print(f"  Fonts: {list(set(fonts)) if fonts else 'default'}")
        print(f"  Sizes: {list(set(sizes)) if sizes else 'default'}")
        print(f"  Bold: {list(set(bolds)) if bolds else 'default'}")
        print(f"  Text: {text[:200]}")

# Print all tables
print("\n\n" + "=" * 80)
print("TABLES IN DOCUMENT")
print("=" * 80)
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_text = cell.text.strip()
            if cell_text:
                print(f"  [{r_idx},{c_idx}]: {cell_text[:150]}")

print("\n\n" + "=" * 80)
print("SECTIONS (Headings analysis)")
print("=" * 80)
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style = para.style.name if para.style else "None"
    if text and ("Heading" in style or "Título" in style or "Title" in style or "head" in style.lower() or para.runs and para.runs[0].font.bold):
        print(f"  Para {i}: [{style}] {text}")
