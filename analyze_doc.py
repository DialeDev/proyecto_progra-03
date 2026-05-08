#!/usr/bin/env python3
"""
Paso 1: Analizar la estructura completa del documento.
Ejecutar: python3 analyze_doc.py
"""

from docx import Document

doc = Document("/home/melquisedec/Escritorio/Projects/Progra III/Proyecto Final/Project Document.docx")

print("=" * 80)
print("ESTRUCTURA COMPLETA DEL DOCUMENTO")
print("=" * 80)

# 1. Todos los párrafos con estilo
print("\n\n--- PÁRRAFOS CON ESTILO ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style = para.style.name if para.style else "None"
    if text:
        # Font info
        fonts = set()
        sizes = set()
        for run in para.runs:
            if run.font.name:
                fonts.add(run.font.name)
            if run.font.size:
                sizes.add(str(run.font.size))
        
        print(f"\nP{i:03d} | Style: {style:30s} | Font: {fonts} | Size: {sizes}")
        print(f"     Text: {text[:300]}")

# 2. Tablas
print("\n\n--- TABLAS ---")
for t_idx, table in enumerate(doc.tables):
    print(f"\nTabla {t_idx}: {len(table.rows)} filas x {len(table.columns)} columnas")
    for r_idx, row in enumerate(table.rows):
        row_data = []
        for c_idx, cell in enumerate(row.cells):
            row_data.append(cell.text.strip()[:100])
        print(f"  Fila {r_idx}: {row_data}")

# 3. Identificar secciones/encabezados
print("\n\n--- ENCABEZADOS (HEADINGS) ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style = para.style.name if para.style else ""
    if text and ("heading" in style.lower() or "Heading" in style or "títu" in style.lower() or "Título" in style):
        print(f"P{i:03d} | {style:30s} | {text}")

print("\n\n--- POSIBLES ENCABEZADOS (texto en negrita) ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text and para.runs and para.runs[0].font.bold and para.runs[0].font.size and para.runs[0].font.size >= 200000:
        style = para.style.name if para.style else "None"
        print(f"P{i:03d} | {style:30s} | {text}")
